"""İndirme doğrulaması; gerçek kayıtları değiştirmez."""
import io
from contextlib import closing
import sqlite3
import sys
import tempfile
from pathlib import Path
from datetime import date, timedelta
from openpyxl import load_workbook
from depo import _SEMA
from excel_rapor import olustur


def test_export():
    with tempfile.TemporaryDirectory() as tmp:
        dbfile = str(Path(tmp) / 'test.db')
        with closing(sqlite3.connect(dbfile)) as db:
            db.executescript(_SEMA)
            rows = [('2026-09-10', 12, 'browser.exe', 'Tarayıcı', '=1+1',
                     'https://example.test/' + str(i) + ('x' * 420 if i == 0 else ''),
                     'example.test', 'gorunur', 120., 60.)
                    for i in range(601)]
            rows.append(('2026-09-11', 13, 'other.exe', 'Diğer', 'Hariç', '', '', 'simge', 999., 0.))
            db.executemany('INSERT INTO kayit VALUES (?,?,?,?,?,?,?,?,?,?)', rows)
            db.execute('INSERT INTO ozet VALUES (?,?,?,?)', ('2026-09-10', 90000., 80000., 10000.))
            db.commit()
        book = load_workbook(io.BytesIO(olustur(dbfile, '2026-09-10', '2026-09-10')))
        assert book['Uygulama içi ayrıntılar'].max_row == 605, '601 farklı adres eksiksiz olmalı'
        assert book['Gün ve saat ayrıntıları'].max_row == 605
        assert book['Uygulama toplamları']['C5'].value == timedelta(seconds=601*60)
        assert book['Uygulama içi ayrıntılar']['C5'].data_type == 's', 'Başlık formül olmamalı'
        assert book['Günlük özet']['B5'].value == timedelta(seconds=90000)
        grouped = book.worksheets[0]
        assert grouped.title == 'Uygulamalar ve sayfalar'
        assert grouped.max_row == 606
        assert grouped['B5'].value == 'Tarayıcı'
        assert grouped['G5'].value == 'browser.exe'
        assert grouped['B6'].value == '=1+1'
        assert grouped['G6'].value.startswith('https://example.test/0')
        assert grouped['B6'].data_type == 's'
        assert grouped['C5'].value == sum((r[2].value for r in grouped.iter_rows(min_row=6)), timedelta())
        assert grouped.row_dimensions[6].outlineLevel == 1
        assert not grouped.row_dimensions[6].hidden
        assert grouped['B6'].alignment.vertical == grouped['C6'].alignment.vertical == 'center'
        assert grouped['C6'].value == timedelta(seconds=60)
        assert grouped.row_dimensions[6].height == 32
        assert grouped.row_dimensions[5].height == 36
        assert grouped['C4'].alignment.vertical == 'center'
        assert book.worksheets[0].freeze_panes == 'C5'
        empty = load_workbook(io.BytesIO(olustur(dbfile, '2020-01-01', '2020-01-01')))
        assert all(s.max_row == 4 for s in empty)
        for start, end in [('bad', 'bad'), ('2026-09-11', '2026-09-10')]:
            try:
                olustur(dbfile, start, end)
            except ValueError:
                pass
            else:
                raise AssertionError('Geçersiz tarih reddedilmeli')
    print('PASS: 601 adres, tarih sınırı, süreler, metin güvenliği, boş aralık')


if __name__ == '__main__':
    test_export()
    today = date.today()
    live_source = str(Path(sys.argv[1]).resolve()) if len(sys.argv) > 1 else str(Path(__file__).with_name('veri.db'))
    snapshot_dir = tempfile.TemporaryDirectory()
    source = str(Path(snapshot_dir.name) / 'snapshot.db')
    with closing(sqlite3.connect(live_source)) as canlı, closing(sqlite3.connect(source)) as kopya:
        canlı.backup(kopya)
    ranges = [(today, today), (today-timedelta(days=1), today-timedelta(days=1)),
              (today-timedelta(days=today.weekday()), today), (today.replace(day=1), today)]
    for start, end in ranges:
        book = load_workbook(io.BytesIO(olustur(source, start.isoformat(), end.isoformat())))
        with closing(sqlite3.connect('file:' + source + '?mode=ro', uri=True)) as db:
            expected = db.execute('SELECT COUNT(*), COALESCE(SUM(aktif_saniye),0) FROM kayit WHERE gun BETWEEN ? AND ?', (start.isoformat(), end.isoformat())).fetchone()
        assert book['Gün ve saat ayrıntıları'].max_row-4 == expected[0]
        total = sum((r[2].value.total_seconds() for r in book['Uygulama toplamları'].iter_rows(min_row=5)), 0)
        assert abs(total-expected[1]) < .1
        print('PASS:', start, end, expected[0], 'kayıt')

    import threading
    import base64
    import hashlib
    import json
    import re
    import urllib.error
    import urllib.request
    from types import SimpleNamespace
    from http.server import ThreadingHTTPServer
    import arayuz
    from kimlik import Kimlik
    authdir = tempfile.TemporaryDirectory()
    arayuz._kimlik = Kimlik(str(Path(authdir.name) / 'kimlik.json'))
    sekme_istekleri = []
    arayuz._takipci = SimpleNamespace(
        depo=SimpleNamespace(dosya=source), diske_yaz=lambda: None,
        zaman=SimpleNamespace(token='yalnizca-test', receive=sekme_istekleri.append),
        _hata_yaz=lambda hata: None)
    server = ThreadingHTTPServer(('127.0.0.1', 0), arayuz._Islemci)
    worker = threading.Thread(target=server.serve_forever, daemon=True)
    worker.start()
    try:
        address = 'http://127.0.0.1:' + str(server.server_port)
        def script_izinli_mi(response):
            govde = response.read().decode('utf-8')
            script = re.search(r'<script>(.*?)</script>', govde, re.S).group(1)
            ozet = base64.b64encode(hashlib.sha256(script.encode('utf-8')).digest()).decode('ascii')
            assert "'sha256-" + ozet + "'" in response.headers['Content-Security-Policy']
            return govde
        with urllib.request.urlopen(address) as response:
            assert 'Şifre oluştur' in script_izinli_mi(response)
            assert response.headers['X-Frame-Options'] == 'DENY'
        try:
            urllib.request.urlopen(urllib.request.Request(address, headers={'Host':'baska.example'}))
        except urllib.error.HTTPError as error:
            assert error.code == 421
        else:
            raise AssertionError('Yabancı sunucu adı reddedilmeli')
        try:
            urllib.request.urlopen(address+'/api/excel?bas=2026-09-07&bit=2026-09-13')
        except urllib.error.HTTPError as error:
            assert error.code == 401
        else:
            raise AssertionError('Şifresiz rapor reddedilmeli')
        request = urllib.request.Request(address+'/auth/kur',
            data=json.dumps({'sifre':'guclu-deneme-sifresi'}).encode('utf-8'),
            headers={'Content-Type':'application/json','X-Ekran-Form':'1',
                     'Origin':'https://baska.example'})
        try:
            urllib.request.urlopen(request)
        except urllib.error.HTTPError as error:
            assert error.code == 403
        else:
            raise AssertionError('Yabancı sayfadan ilk kurulum reddedilmeli')
        request = urllib.request.Request(address+'/auth/kur',
            data=json.dumps({'sifre':'guclu-deneme-sifresi'}).encode('utf-8'),
            headers={'Content-Type':'application/json','X-Ekran-Form':'1'})
        with urllib.request.urlopen(request) as response:
            cookie = response.headers['Set-Cookie'].split(';')[0]
        sekme_istegi = urllib.request.Request(address+'/api/sekme', data=b'{}',
            headers={'Content-Type':'application/json','X-Ekran-Token':'yalnizca-test',
                     'Origin':'chrome-extension://deneme'})
        with urllib.request.urlopen(sekme_istegi) as response:
            assert response.status == 200
        assert sekme_istekleri == [{}]
        with urllib.request.urlopen(urllib.request.Request(address, headers={'Cookie':cookie})) as response:
            assert 'id="excelIndir"' in script_izinli_mi(response)
        with urllib.request.urlopen(urllib.request.Request(address+'/api/excel?bas=2026-09-07&bit=2026-09-13', headers={'Cookie':cookie})) as response:
            assert response.status == 200
            assert response.headers['Content-Type'].endswith('spreadsheetml.sheet')
            assert 'attachment;' in response.headers['Content-Disposition']
            book = load_workbook(io.BytesIO(response.read()))
            assert len(book.worksheets) == 7
        try:
            urllib.request.urlopen(urllib.request.Request(address+'/api/excel?bas=bad&bit=bad', headers={'Cookie':cookie}))
        except urllib.error.HTTPError as error:
            assert error.code == 400
        else:
            raise AssertionError('Hatalı tarih reddedilmeli')
        print('PASS: düğme, indirme yanıtı, dosya adı, dosyayı açma, hatalı tarih')
    finally:
        server.shutdown()
        server.server_close()
        worker.join()
        authdir.cleanup()
        snapshot_dir.cleanup()
