import io
import os
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta
from openpyxl import load_workbook
from depo import Depo
from zaman import Zaman
from excel_rapor import olustur


class TimelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.file = os.path.join(self.temp.name, 'test.db')
        d = Depo(self.file)
        d.kapat()
        self.z = Zaman(self.file)
        self.t = datetime(2026,9,15,9).timestamp()
        self.window = dict(hwnd=1,pid=11,exe='chrome.exe',ad='Google Chrome',baslik='Video - Google Chrome',state='Ekranda görünür',on_planda=True)

    def tearDown(self):
        self.z.close()
        self.temp.cleanup()

    def send(self, seconds=0, active=1, closed=(), tabs=(1,2), seq=None):
        self.z.receive(dict(session='session1',seq=seq or seconds+1,browser='chrome.exe',closed=list(closed),tabs=[dict(id=i,window=1,title='Video' if i==1 else 'Belge',url='https://site.test/'+str(i),active=i==active,focused=True) for i in tabs]),now=self.t+seconds)

    def test_tab_switch_cover_close(self):
        self.z.observe([self.window],now=self.t)
        self.send()
        self.send(2,active=2)
        self.z.observe([dict(self.window,baslik='Belge - Google Chrome',state='Arka planda — üstü kapalı veya ekran dışında',on_planda=False)],now=self.t+4)
        self.send(6,active=2,closed=(1,),tabs=(2,))
        rows=list(self.z.db.execute("SELECT * FROM zaman_aralik WHERE kimlik='s:session1:1' ORDER BY id"))
        self.assertEqual(rows[0]['durum'],'Ekranda görünür')
        self.assertIn('başka sekme',rows[1]['durum'])
        self.assertEqual(rows[-1]['bitis_nedeni'],'Sekme kapandı')
        rows=list(self.z.db.execute("SELECT * FROM zaman_aralik WHERE kimlik='s:session1:2' ORDER BY id"))
        self.assertIn('üstü kapalı',rows[-1]['durum'])
        self.assertIsNone(rows[-1]['bitis'])

    def test_unchanged_does_not_grow(self):
        self.z.observe([self.window],now=self.t)
        for i in range(2,62,2):
            self.z.observe([self.window],now=self.t+i)
        self.z.flush()
        self.assertEqual(self.z.db.execute('SELECT COUNT(*) FROM zaman_aralik').fetchone()[0],1)
        r=self.z.db.execute('SELECT * FROM zaman_aralik').fetchone()
        self.assertEqual(datetime.fromisoformat(r['son_gorulme'])-datetime.fromisoformat(r['baslangic']),timedelta(seconds=60))

    def test_gap_not_fabricated(self):
        self.z.observe([self.window],now=self.t)
        self.z.observe([self.window],now=self.t+2)
        self.z.observe([self.window],now=self.t+100)
        rows=list(self.z.db.execute('SELECT * FROM zaman_aralik ORDER BY id'))
        self.assertIn('Ölçüm kesildi',rows[0]['bitis_nedeni'])
        self.assertEqual(datetime.fromisoformat(rows[0]['bitis']),datetime.fromtimestamp(self.t+2))

    def test_duplicate_and_stale_browser(self):
        self.z.observe([self.window],now=self.t)
        self.send()
        n=self.z.db.execute('SELECT COUNT(*) FROM zaman_aralik').fetchone()[0]
        self.send(2,seq=1)
        self.assertEqual(n,self.z.db.execute('SELECT COUNT(*) FROM zaman_aralik').fetchone()[0])
        for i in range(2,100,2): self.z.observe([self.window],now=self.t+i)
        rows=list(self.z.db.execute("SELECT * FROM zaman_aralik WHERE kaynak='Sekme'"))
        self.assertTrue(all('bağlantısı kesildi' in r['bitis_nedeni'] for r in rows))

    def test_export_real_dates_and_open_end(self):
        self.z.observe([self.window],now=self.t)
        self.send()
        self.z.observe([self.window],now=self.t+2)
        self.z.flush()
        b=load_workbook(io.BytesIO(olustur(self.file,'2026-09-15','2026-09-15')))
        ws=b['Ayrıntılı zaman çizelgesi']
        self.assertEqual(ws.max_row,7)
        self.assertIsInstance(ws['A5'].value,datetime)
        self.assertIsNone(ws['B5'].value)
        self.assertEqual(ws['C5'].value,timedelta(seconds=2))
        self.assertIn('devam ediyor',ws['H5'].value)
        tabs=b['Sekme saatleri']
        self.assertEqual(tabs.max_row,6)
        self.assertEqual(tabs['A4'].value,'Başlangıç')
        self.assertEqual(tabs['C4'].value,'Doğrulanan süre')
        self.assertEqual(tabs['L5'].value,'Sekme')

    def test_offscreen(self):
        import pencere
        self.assertEqual(pencere.gorunur_yuzde((200,0,300,100),[],[(0,0,100,100)]),0)
        self.assertEqual(pencere.gorunur_yuzde((0,0,100,100),[(0,0,100,100)],[(0,0,100,100)]),0)
        self.assertEqual(pencere.gorunur_yuzde((0,0,200,100),[],[(0,0,100,100)]),50)


if __name__=='__main__': unittest.main()
