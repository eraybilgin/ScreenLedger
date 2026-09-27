"""Seçilen aralığın tüm kayıtlarını, yerel Excel indirmesi için hazırlar."""
from collections import defaultdict
from contextlib import closing
from datetime import date, datetime, timedelta
from io import BytesIO
import sqlite3

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def olustur(dosya, bas, bit):
    bas, bit = date.fromisoformat(bas), date.fromisoformat(bit)
    if bas > bit:
        raise ValueError('Başlangıç tarihi bitiş tarihinden sonra olamaz.')
    # Tek okuma işlemi: kayıtlar ve günlük özet aynı anı temsil eder.
    with closing(sqlite3.connect('file:' + dosya + '?mode=ro', uri=True)) as db:
        db.row_factory = sqlite3.Row
        db.execute('BEGIN')
        kayitlar = [dict(r) for r in db.execute(
            'SELECT * FROM kayit WHERE gun BETWEEN ? AND ? ORDER BY exe, baslik, url, gun, saat, durum',
            (bas.isoformat(), bit.isoformat()))]
        gunler = [dict(r) for r in db.execute(
            'SELECT * FROM ozet WHERE gun BETWEEN ? AND ? ORDER BY gun',
            (bas.isoformat(), bit.isoformat()))]
        araliklar = []
        if db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='zaman_aralik'").fetchone():
            araliklar = [dict(r) for r in db.execute(
                'SELECT * FROM zaman_aralik WHERE baslangic < ? AND son_gorulme >= ? ORDER BY uygulama,kimlik,baslangic,id',
                ((bit+timedelta(days=1)).isoformat(), bas.isoformat()))]

    kitap = Workbook()
    kitap.remove(kitap.active)
    kitap.properties.title = 'Ekran Takip kullanım raporu'
    kitap.properties.creator = 'Ekran Takip'

    def sayfa(ad, basliklar, satirlar, genislikler, sure_sutunlari=()):
        ws = kitap.create_sheet(ad)
        ws.append([ad + ' — Ekran Takip'])
        ws.append([f'{bas:%d.%m.%Y} – {bit:%d.%m.%Y} | Oluşturulma: {datetime.now():%d.%m.%Y %H:%M:%S}'])
        ws.append(['Süreler saat:dakika:saniye biçimindedir. Yalnızca kaydedilmiş sayfa ve adresler yer alır.'])
        ws.append(basliklar)
        for row in satirlar:
            ws.append(row)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    # Sayfa başlıkları ve adresler hiçbir zaman formül çalıştırmaz.
                    cell.value = ILLEGAL_CHARACTERS_RE.sub('', cell.value)
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i in (1, 2, 3):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(basliklar))
        ws['A1'].font = Font(size=17, bold=True, color='1D4ED8')
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 23
        ws.row_dimensions[3].height = 30
        ws.row_dimensions[4].height = 34
        for cell in ws[4]:
            cell.fill = PatternFill('solid', fgColor='2563EB')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in ws.iter_rows(min_row=5):
            ws.row_dimensions[row[0].row].height = 44
            for cell in row:
                if cell.row % 2:
                    cell.fill = PatternFill('solid', fgColor='EFF6FF')
                if cell.column in sure_sutunlari:
                    cell.number_format = '[h]:mm:ss'
                elif isinstance(cell.value, datetime):
                    cell.number_format = 'dd.mm.yyyy hh:mm:ss'
                elif isinstance(cell.value, date):
                    cell.number_format = 'dd.mm.yyyy'
        for i, width in enumerate(genislikler, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = 'C5'
        ws.auto_filter.ref = f'A4:{get_column_letter(len(basliklar))}{max(4, ws.max_row)}'
        ws.sheet_view.showGridLines = False
        return ws

    toplam = defaultdict(lambda: [0., 0., 0., 0.])
    detay = defaultdict(lambda: [0., 0., 0., 0.])
    adlar = {}
    durum_sutunu = {'gorunur': 1, 'ustu_kapali': 2, 'simge': 3}
    for r in kayitlar:
        adlar[r['exe']] = r['ad']
        for v in (toplam[r['exe']], detay[(r['exe'], r['baslik'], r['url'], r['alan'])]):
            v[0] += r['aktif_saniye']
            v[durum_sutunu[r['durum']]] += r['saniye']
    sureler = ['Aktif kullanım', 'Ekranda görünür', 'Üstü kapalı', 'Simge durumunda']
    def zamanlar(v):
        return [timedelta(seconds=s) for s in v]
    uygulamalar = sorted(toplam, key=lambda e: (-toplam[e][0], e))
    sayfa('Uygulama toplamları', ['Uygulama', 'Program dosyası'] + sureler,
          [[adlar[e], e] + zamanlar(toplam[e]) for e in uygulamalar],
          [32, 28, 23, 23, 23, 23], (3, 4, 5, 6))
    sira = {e: i for i, e in enumerate(uygulamalar)}
    # İlk görünüm: uygulama toplamı ve hemen altında bütün sayfaları.
    gruplar = defaultdict(list)
    for k, v in detay.items():
        gruplar[k[0]].append((k, v))
    rapor_satirlari, toplam_satirlari = [], set()
    for i, e in enumerate(uygulamalar, 1):
        toplam_satirlari.add(len(rapor_satirlari) + 5)
        rapor_satirlari.append([i, adlar[e]] + zamanlar(toplam[e]) + [e])
        for k, v in sorted(gruplar[e], key=lambda kv: (-kv[1][0], -kv[1][1], kv[0])):
            rapor_satirlari.append(['', k[1]] + zamanlar(v) + [k[2]])
    rapor = sayfa('Uygulamalar ve sayfalar',
                  ['#', 'Uygulama / sayfa'] + sureler + ['Adres / program dosyası'],
                  rapor_satirlari, [6, 90, 23, 23, 23, 23, 70], (3, 4, 5, 6))
    # Gruplu görünümde sıralama üst ve alt satırları ayırmamalı.
    rapor.auto_filter.ref = None
    rapor.freeze_panes = 'C5'
    rapor.sheet_properties.outlinePr.summaryBelow = False
    for row in rapor:
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='161920')
            cell.font = Font(name='Calibri', size=11, color='AFC1D3')
            cell.alignment = Alignment(
                horizontal='right' if 3 <= cell.column <= 6 else 'left',
                vertical='center', wrap_text=False)
    rapor['A1'].font = Font(name='Calibri', size=17, bold=True, color='60A5FA')
    for cell in rapor[4]:
        cell.fill = PatternFill('solid', fgColor='24344B')
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in rapor.iter_rows(min_row=5):
        no = row[0].row
        ana = no in toplam_satirlari
        if ana:
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='253750')
                cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
            rapor.row_dimensions[no].height = 36
        else:
            rapor.row_dimensions[no].outlineLevel = 1
            row[1].alignment = Alignment(vertical='center', wrap_text=False, indent=1)
            row[6].font = Font(name='Calibri', size=10, color='8FA5BD')
            rapor.row_dimensions[no].height = 32
    kitap.move_sheet(rapor, offset=-1)
    sayfa('Uygulama içi ayrıntılar',
          ['Uygulama', 'Program dosyası', 'Sayfa / pencere başlığı', 'Tam adres', 'Site'] + sureler,
          [[adlar[k[0]], k[0], k[1], k[2], k[3]] + zamanlar(v)
           for k, v in sorted(detay.items(), key=lambda kv: (sira[kv[0][0]], -kv[1][0], kv[0]))],
          [30, 24, 65, 75, 30, 23, 23, 23, 23], (6, 7, 8, 9))
    durum_adlari = {'gorunur': 'Ekranda görünür', 'ustu_kapali': 'Üstü kapalı', 'simge': 'Simge durumunda'}
    sayfa('Gün ve saat ayrıntıları',
          ['Tarih', 'Saat aralığı', 'Uygulama', 'Program dosyası', 'Sayfa / pencere başlığı', 'Tam adres', 'Site', 'Pencere durumu', 'Durum süresi', 'Aktif kullanım'],
          [[date.fromisoformat(r['gun']), f"{r['saat']:02d}:00–{r['saat']:02d}:59", r['ad'], r['exe'], r['baslik'], r['url'], r['alan'], durum_adlari[r['durum']], timedelta(seconds=r['saniye']), timedelta(seconds=r['aktif_saniye'])] for r in kayitlar],
          [16, 20, 30, 24, 65, 75, 30, 24, 23, 23], (9, 10))
    sayfa('Günlük özet', ['Tarih', 'Bilgisayar açık', 'Başındaydın', 'Boşta geçti'],
          [[date.fromisoformat(r['gun'])] + zamanlar([r['acik_saniye'], r['etkin_saniye'], r['bosta_saniye']]) for r in gunler],
          [20, 28, 28, 28], (2, 3, 4))
    alt = datetime.combine(bas, datetime.min.time())
    ust = datetime.combine(bit+timedelta(days=1), datetime.min.time())
    hareketler = []
    for r in araliklar:
        ilk = max(alt, datetime.fromisoformat(r['baslangic']))
        son = min(ust, datetime.fromisoformat(r['son_gorulme']))
        if son < ilk:
            continue
        acik = r['bitis'] is None
        notu = r['bitis_nedeni']
        if acik:
            notu = 'Son gözlemde devam ediyor; bitiş bilinmiyor'
        if datetime.fromisoformat(r['son_gorulme']) >= ust:
            notu = 'Seçili aralıktan sonra devam ediyor'
        hareketler.append([r['kaynak'],r['uygulama'],r['baslik'],r['adres'],r['kimlik'],ilk,
                          None if acik and son < ust else son,son,son-ilk,r['durum'],r['kullanim'],notu])
    def okunur_satir(r):
        # Saat ve süreyi başa al: geniş adres/kimlik sütunları artık onları
        # ekranın dışına itmez.
        return [r[5],r[6],r[8],r[1],r[2],r[9],r[10],r[11],r[3],r[7],r[4],r[0]]

    sutunlar = ['Başlangıç','Bitiş','Doğrulanan süre','Uygulama',
                'Sekme / pencere','Görünürlük durumu','Kullanıcı durumu',
                'Aralığın bitme nedeni','Tam adres','Son doğrulanan saat',
                'Pencere / sekme kimliği','Kayıt türü']
    genislikler = [25,25,23,27,65,48,25,58,70,25,30,19]
    sekmeler = [r for r in hareketler if r[0].startswith('Sekme')]
    sekmeler.sort(key=lambda r: (r[1].lower(),r[5],r[4]))
    focus = sayfa('Sekme saatleri',sutunlar,[okunur_satir(r) for r in sekmeler],
                  genislikler,(3,))
    focus['A3'] = 'Her satır bir sekmenin durum aralığıdır. Boş bitiş: son gözlemde sürüyor. Doğrulanan süre, başlangıç ile son doğrulanan saat arasındadır.'
    focus.row_dimensions[3].height = 42
    focus.freeze_panes = 'D5'
    kitap.move_sheet(focus, offset=1-kitap.index(focus))

    ws = sayfa('Ayrıntılı zaman çizelgesi',sutunlar,
               [okunur_satir(r) for r in hareketler],genislikler,(3,))
    ws['A3'] = 'Pencere ve sekme satırları aynı zamanı gösterebilir; birlikte toplamayın. Boş bitiş: son gözlemde sürüyor. Doğrulanan süre kesin kapanış saati değildir.'
    ws.row_dimensions[3].height = 42
    ws.freeze_panes = 'D5'
    for sheet in (focus, ws):
        # Uzun başlıklar alta sarkıp sonraki sürenin hizasını bozmasın.
        for row in sheet.iter_rows(min_row=5):
            title = str(row[4].value or '')
            sheet.row_dimensions[row[0].row].height = min(120, max(44, 16 * ((len(title)+59)//60) + 12))
    cikti = BytesIO()
    kitap.save(cikti)
    return cikti.getvalue()
